from dataclasses import dataclass
from typing import Dict, List, Tuple, Optional
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import re, unicodedata

def _norm_text(x) -> str:
    if x is None: return ""
    s = str(x).strip().upper()
    s = "".join(ch for ch in unicodedata.normalize("NFD", s) if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", s)

def _to_number_ar(v) -> float:
    if v is None: return 0.0
    if isinstance(v, bool): return float(int(v))
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "")
    if not s: return 0.0
    if "." in s and "," in s: s = s.replace(".", "").replace(",", ".")
    elif "," in s and "." not in s: s = s.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s): s = s.replace(".", "")
    s = re.sub(r"[^0-9.\-]", "", s)
    try: return float(s)
    except: return 0.0

DAY_ALIASES = {
    "Dom": ["DOM"], "Lun": ["LUN"], "Mar": ["MAR"],
    "Miérc": ["MIERC", "MIER", "MIE"], "Juev": ["JUEV", "JUE"],
    "Vier": ["VIER", "VIE"], "Sáb": ["SAB", "SÁB"]
}

def _detect_day_header_row(ws, max_row=80, max_col=120) -> Optional[int]:
    for r in range(1, min(ws.max_row, max_row) + 1):
        found = set()
        for c in range(1, min(ws.max_column, max_col) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                t = _norm_text(v)
                for lab, aliases in DAY_ALIASES.items():
                    if any(t.startswith(a) for a in aliases): found.add(lab)
        if "Dom" in found and len(found) >= 5: return r
    return None

def _detect_name_col(ws, max_row=80, max_col=120) -> Optional[int]:
    for r in range(1, min(ws.max_row, max_row) + 1):
        for c in range(1, min(ws.max_column, max_col) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and "APELLIDO" in _norm_text(v) and "NOMBRE" in _norm_text(v): return c
    return None

def _detect_day_cols(ws, day_row: int, max_col=200) -> Dict[str, int]:
    cols = {}
    for c in range(1, min(ws.max_column, max_col) + 1):
        v = ws.cell(day_row, c).value
        if isinstance(v, str):
            t = _norm_text(v)
            for lab, aliases in DAY_ALIASES.items():
                if any(t.startswith(a) for a in aliases): cols[lab] = c
    return cols

def _detect_holiday_col(ws, day_row: int, max_col=200) -> Optional[int]:
    for c in range(1, min(ws.max_column, max_col) + 1):
        v = ws.cell(day_row, c).value
        if isinstance(v, str) and "FERIADO" in _norm_text(v): return c
    return None

def _detect_totals_header_row(ws, start_row: int, look_ahead=4, max_col=250) -> int:
    best_row, best_score = start_row, 0
    for r in range(start_row, start_row + look_ahead):
        score = sum(1 for c in range(1, min(ws.max_column, max_col) + 1)
                    if isinstance(ws.cell(r, c).value, str) and any(k in _norm_text(ws.cell(r, c).value) for k in ["$/H","SUB TOTAL","SUBTOTAL","REDONDEO","TOTAL"]))
        if score > best_score: best_score, best_row = score, r
    return best_row

def _detect_totals_cols(ws, totals_row: int, max_col=250) -> Dict[str, int]:
    cols = {}
    for c in range(1, min(ws.max_column, max_col) + 1):
        v = ws.cell(totals_row, c).value
        if isinstance(v, str):
            t = _norm_text(v)
            if "$/H" in t and "LAV" in t: cols["lv"] = c
            elif "$/H" in t and "SAB" in t: cols["sab"] = c
            elif "$/H" in t and ("DOM" in t or "FER" in t): cols["domfer"] = c
            elif "SUB TOTAL" in t or t == "SUBTOTAL": cols["subtotal"] = c
            elif "REDONDEO" in t: cols["redondeo"] = c
    return cols

def _guess_data_start(day_header_row: int) -> int: return day_header_row + 4

@dataclass
class EmpRow:
    name: str
    hours: Dict[str, float]
    totals: Dict[str, float]

def _read_weekly_file(path_xlsx: str) -> Tuple[Dict, List[EmpRow]]:
    wb = load_workbook(path_xlsx, data_only=True)
    for ws in wb.worksheets:
        day_row = _detect_day_header_row(ws)
        if not day_row: continue
        name_col = _detect_name_col(ws)
        if not name_col: raise ValueError(f"No encontré 'APELLIDO Y NOMBRE' en {path_xlsx}")
        day_cols = _detect_day_cols(ws, day_row)
        if "Dom" not in day_cols: raise ValueError(f"No encontré 'Dom' en {path_xlsx}")
        hol_col = _detect_holiday_col(ws, day_row)
        totals_row = _detect_totals_header_row(ws, day_row)
        totals_cols = _detect_totals_cols(ws, totals_row)
        data_start = _guess_data_start(day_row)

        day_dates = {lab: ws.cell(day_row + 1, c).value for lab, c in day_cols.items()}
        holiday_date = ws.cell(day_row + 1, hol_col).value if hol_col else None

        rows, empty_streak = [], 0
        for r in range(data_start, ws.max_row + 1):
            name_v = ws.cell(r, name_col).value
            name = str(name_v).strip() if isinstance(name_v, str) else ""
            if not name:
                empty_streak += 1
                if empty_streak >= 30: break
                continue
            empty_streak = 0
            hours = {lab: _to_number_ar(ws.cell(r, day_cols[lab]).value) for lab in ["Dom","Lun","Mar","Miérc","Juev","Vier","Sáb"] if day_cols.get(lab)}
            hours["Feriado"] = _to_number_ar(ws.cell(r, hol_col).value) if hol_col else 0.0
            totals = {k: _to_number_ar(ws.cell(r, totals_cols.get(k, 0)).value) if totals_cols.get(k) else 0.0 for k in ["lv","sab","domfer","subtotal","redondeo","total"]}
            if totals["total"] == 0 and (totals["subtotal"] != 0 or totals["redondeo"] != 0):
                totals["total"] = totals["subtotal"] + totals["redondeo"]
            rows.append(EmpRow(name=name, hours=hours, totals=totals))
        return {"day_dates": day_dates, "holiday_date": holiday_date}, rows
    raise ValueError(f"No encontré hoja con días en {path_xlsx}")

def build_rrhh_print_workbook(path_oeste: str, path_consultora: str, template_print_path: str, out_path: str) -> None:
    meta_o, rows_o = _read_weekly_file(path_oeste)
    meta_c, rows_c = _read_weekly_file(path_consultora)
    meta = meta_o if meta_o.get("day_dates") else meta_c

    merged: Dict[str, EmpRow] = {}
    def add_rows(lst):
        for er in lst:
            k = _norm_text(er.name)
            if k not in merged: merged[k] = er
            else:
                cur = merged[k]
                for d, v in er.hours.items(): cur.hours[d] = _to_number_ar(cur.hours.get(d, 0)) + _to_number_ar(v)
                for t, v in er.totals.items(): cur.totals[t] = _to_number_ar(cur.totals.get(t, 0)) + _to_number_ar(v)

    add_rows(rows_o); add_rows(rows_c)
    final_rows = sorted(merged.values(), key=lambda x: _norm_text(x.name))

    wb = load_workbook(template_print_path)
    ws = wb["impresion"] if "impresion" in wb.sheetnames else wb.active

    day_row = _detect_day_header_row(ws)
    name_col = _detect_name_col(ws)
    num_col = None
    for c in range(1, min(ws.max_column, 40)+1):
        v = ws.cell(day_row, c).value
        if isinstance(v, str) and _norm_text(v).startswith("NUM"): num_col = c; break
    if not name_col or not num_col: raise ValueError("La plantilla debe tener columnas 'NUM' y 'APELLIDO Y NOMBRE'.")

    day_cols = _detect_day_cols(ws, day_row)
    hol_col = _detect_holiday_col(ws, day_row)
    totals_row = _detect_totals_header_row(ws, day_row)
    totals_cols = _detect_totals_cols(ws, totals_row)
    data_start = _guess_data_start(day_row)
    dates_row = day_row + 1

    for lab, dt in meta.get("day_dates", {}).items():
        c = day_cols.get(lab)
        if c: ws.cell(dates_row, c).value = dt
    if hol_col: ws.cell(dates_row, hol_col).value = meta.get("holiday_date")

    relevant_cols = {num_col, name_col}
    relevant_cols.update(day_cols.values())
    if hol_col: relevant_cols.add(hol_col)
    relevant_cols.update(totals_cols.values())

    last_clear = max(ws.max_row, data_start + 400)
    for r in range(data_start, last_clear + 1):
        for c in relevant_cols: ws.cell(r, c).value = None

    for idx, er in enumerate(final_rows, start=1):
        r = data_start + idx - 1
        ws.cell(r, num_col).value = idx
        ws.cell(r, name_col).value = er.name
        for lab in ["Dom","Lun","Mar","Miérc","Juev","Vier","Sáb"]:
            c = day_cols.get(lab)
            if c: ws.cell(r, c).value = _to_number_ar(er.hours.get(lab, 0))
        if hol_col: ws.cell(r, hol_col).value = _to_number_ar(er.hours.get("Feriado", 0))
        for k in ["lv","sab","domfer","subtotal","redondeo","total"]:
            c = totals_cols.get(k)
            if c: ws.cell(r, c).value = _to_number_ar(er.totals.get(k, 0))

    if final_rows:
        last_row = data_start + len(final_rows) - 1
        ws.print_area = f"{get_column_letter(min(relevant_cols))}1:{get_column_letter(max(relevant_cols))}{last_row}"
    wb.save(out_path)